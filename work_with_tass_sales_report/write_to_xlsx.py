from openpyxl import load_workbook


def write_to_main_file(photos, main_report, report_date):  # записываю информацию в главный файл
    try:
        # Открываем файл Excel
        wb = load_workbook(filename=main_report, read_only=False)

        # Создаем новый лист
        sheet_name = " ".join(report_date)
        if sheet_name in wb.sheetnames:
            raise ValueError(f"Лист '{sheet_name}' уже существует в файле {main_report}")

        ws_month_number = wb.create_sheet(sheet_name, 0)

        # Заполняем заголовки
        ws_month_number.cell(row=1, column=1).value = "photo_id"
        ws_month_number.cell(row=1, column=2).value = "income"
        ws_month_number.cell(row=1, column=3).value = "sold times"

        # Заполняем данные
        for idx, (photo_id, incomes) in enumerate(photos.items(), start=2):
            ws_month_number.cell(row=idx, column=1).value = photo_id
            ws_month_number.cell(row=idx, column=2).value = sum(incomes)  # суммирую доход по снимкам
            ws_month_number.cell(row=idx, column=3).value = len(incomes)

        # Сохраняем изменения
        wb.save(main_report)
        print("Информация записана.")

    except Exception as e:
        print(f"Ошибка при записи в файл: {e}")

    finally:
        # Гарантируем освобождение ресурсов
        wb.close()